"""Word document report generator."""
import logging
from datetime import datetime
from pathlib import Path
from typing import Any, Dict

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

logger = logging.getLogger(__name__)


class WordGenerator:
    """Generates Word document reports from prospect data."""

    def __init__(self, output_dir: str = "./reports"):
        """Initialize Word generator.

        Args:
            output_dir: Directory to save Word files
        """
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate(self, report_data: Dict[str, Any]) -> str:
        """Generate Word document report.

        Args:
            report_data: Complete report data dictionary

        Returns:
            Path to generated Word file
        """
        logger.info("Generating Word document report")

        # Create filename
        prospect_name = report_data.get("prospect_name", "prospect").replace(" ", "_")
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{prospect_name}_report_{timestamp}.docx"
        filepath = self.output_dir / filename

        # Create document
        doc = Document()

        # Title page
        self._create_title_page(doc, report_data)
        doc.add_page_break()

        # Executive summary
        self._create_executive_summary(doc, report_data)

        # Quick facts
        self._create_quick_facts(doc, report_data)

        # Talking points
        self._create_talking_points(doc, report_data)

        # Conversation starters
        self._create_conversation_starters(doc, report_data)

        # Personality analysis
        self._create_personality_analysis(doc, report_data)

        # Recommended approach
        self._create_recommended_approach(doc, report_data)

        # Company information
        self._create_company_info(doc, report_data)

        # News and events
        self._create_news_section(doc, report_data)

        # Save document
        doc.save(str(filepath))

        logger.info(f"Word document generated: {filepath}")
        return str(filepath)

    def _create_title_page(self, doc: Document, report_data: Dict[str, Any]):
        """Create title page."""
        # Title
        title = doc.add_heading("PROSPECT INTELLIGENCE REPORT", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Add some space
        doc.add_paragraph()

        # Prospect info
        prospect_name = report_data.get("prospect_name", "Unknown")
        company_name = report_data.get("company_name", "Unknown")

        info = doc.add_paragraph()
        info.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = info.add_run(f"{prospect_name}\n{company_name}")
        run.font.size = Pt(14)
        run.font.bold = True

        # Date
        doc.add_paragraph()
        date_para = doc.add_paragraph(
            f"Generated on {datetime.now().strftime('%B %d, %Y')}"
        )
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Branding
        doc.add_paragraph()
        branding = doc.add_paragraph("Powered by YUSEARCH - AI-Powered Prospect Research")
        branding.alignment = WD_ALIGN_PARAGRAPH.CENTER
        branding.runs[0].font.italic = True

    def _create_executive_summary(self, doc: Document, report_data: Dict[str, Any]):
        """Create executive summary section."""
        doc.add_heading("Executive Summary", level=1)

        summary = report_data.get("executive_summary", "No summary available")
        doc.add_paragraph(summary)

    def _create_quick_facts(self, doc: Document, report_data: Dict[str, Any]):
        """Create quick facts section."""
        doc.add_heading("Quick Facts", level=1)

        facts = report_data.get("quick_facts", {})

        if facts:
            table = doc.add_table(rows=len(facts), cols=2)
            table.style = "Light Grid Accent 1"

            for i, (key, value) in enumerate(facts.items()):
                row = table.rows[i]
                row.cells[0].text = key
                row.cells[1].text = value

                # Bold the first column
                row.cells[0].paragraphs[0].runs[0].font.bold = True
        else:
            doc.add_paragraph("No quick facts available")

    def _create_talking_points(self, doc: Document, report_data: Dict[str, Any]):
        """Create talking points section."""
        doc.add_heading("Perfect Talking Points", level=1)

        talking_points = report_data.get("talking_points", [])

        if talking_points:
            for i, point in enumerate(talking_points, 1):
                para = doc.add_paragraph(f"{i}. {point}")
                para.style = "List Number"
        else:
            doc.add_paragraph("No talking points available")

    def _create_conversation_starters(self, doc: Document, report_data: Dict[str, Any]):
        """Create conversation starters section."""
        doc.add_heading("Conversation Starters", level=1)

        starters = report_data.get("conversation_starters", [])

        if starters:
            for starter in starters:
                para = doc.add_paragraph(starter, style="List Bullet")
        else:
            doc.add_paragraph("No conversation starters available")

    def _create_personality_analysis(self, doc: Document, report_data: Dict[str, Any]):
        """Create personality analysis section."""
        doc.add_heading("Personality Insights", level=1)

        analysis = report_data.get("personality_analysis", "No analysis available")

        # Split into paragraphs
        for para_text in analysis.split("\n\n"):
            if para_text.strip():
                doc.add_paragraph(para_text.strip())

    def _create_recommended_approach(self, doc: Document, report_data: Dict[str, Any]):
        """Create recommended approach section."""
        doc.add_heading("Recommended Approach", level=1)

        approach = report_data.get("recommended_approach", "No recommendations available")

        for para_text in approach.split("\n\n"):
            if para_text.strip():
                doc.add_paragraph(para_text.strip())

    def _create_company_info(self, doc: Document, report_data: Dict[str, Any]):
        """Create company information section."""
        doc.add_heading("Company Background", level=1)

        company_data = report_data.get("company_data", {})

        if company_data.get("description"):
            doc.add_paragraph(company_data["description"])

        if company_data.get("website"):
            para = doc.add_paragraph()
            para.add_run("Website: ").bold = True
            para.add_run(company_data["website"])

        if company_data.get("industry"):
            para = doc.add_paragraph()
            para.add_run("Industry: ").bold = True
            para.add_run(company_data["industry"])

        if company_data.get("size"):
            para = doc.add_paragraph()
            para.add_run("Company Size: ").bold = True
            para.add_run(company_data["size"])

        if company_data.get("funding"):
            para = doc.add_paragraph()
            para.add_run("Funding: ").bold = True
            para.add_run(company_data["funding"])

    def _create_news_section(self, doc: Document, report_data: Dict[str, Any]):
        """Create recent news section."""
        doc.add_heading("Recent News & Events", level=1)

        news_data = report_data.get("news_data", {})
        company_news = news_data.get("company_news", [])

        if company_news:
            for news in company_news[:5]:  # Top 5 news items
                title = news.get("title", "")
                snippet = news.get("snippet", "")

                # Title
                para = doc.add_paragraph()
                run = para.add_run(title)
                run.bold = True

                # Snippet
                doc.add_paragraph(snippet)

                # Source
                source = news.get("source", "")
                if source:
                    para = doc.add_paragraph()
                    run = para.add_run(f"Source: {source}")
                    run.font.italic = True
                    run.font.size = Pt(9)

                # Add space between news items
                doc.add_paragraph()
        else:
            doc.add_paragraph("No recent news available")
